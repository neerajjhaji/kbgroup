import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
DARK_NAVY = RGBColor(11, 25, 44)       # #0B192C Tata AIA Premium Navy
TATA_BLUE = RGBColor(0, 102, 204)      # #0066CC Tata Blue
ACCENT_CYAN = RGBColor(0, 180, 216)    # #00B4D8 Cloud & AI Cyan
CREAM = RGBColor(248, 249, 250)        # #F8F9FA Off-White Ground
CARD_BG = RGBColor(255, 255, 255)      # White Card
TEXT_DARK = RGBColor(33, 37, 41)
TEXT_MUTED = RGBColor(108, 117, 125)
WHITE = RGBColor(255, 255, 255)
GOLD = RGBColor(212, 175, 55)
GREEN = RGBColor(40, 167, 69)

def set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def add_header(slide, title_text, category_text="TATA AIA LIFE INSURANCE • AVP CLOUD AI OPS ROADMAP"):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = TATA_BLUE
    p0.space_after = Pt(2)

    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = DARK_NAVY

def create_card_slide(prs, title, card1_data, card2_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, CREAM)
    add_header(slide, title)

    top_pos = Inches(1.4)
    card_width = Inches(5.6)
    card_height = Inches(5.4)

    # Left Card
    c1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, card_width, card_height)
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = TATA_BLUE
    c1.line.width = Pt(1.5)

    tf1 = c1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = card1_data['title']
    p1.font.bold = True
    p1.font.size = Pt(15)
    p1.font.color.rgb = TATA_BLUE
    p1.space_after = Pt(10)

    for item in card1_data['items']:
        p = tf1.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # Right Card
    c2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), top_pos, card_width, card_height)
    c2.fill.solid()
    c2.fill.fore_color.rgb = DARK_NAVY
    c2.line.color.rgb = ACCENT_CYAN
    c2.line.width = Pt(1.5)

    tf2 = c2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = card2_data['title']
    p2.font.bold = True
    p2.font.size = Pt(15)
    p2.font.color.rgb = ACCENT_CYAN
    p2.space_after = Pt(10)

    for item in card2_data['items']:
        p = tf2.add_paragraph()
        p.text = f"➔ {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = WHITE
        p.space_after = Pt(6)

# SLIDE 1: Title Slide
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide1, DARK_NAVY)

bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.08), Inches(4.0))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT_CYAN
bar.line.fill.background()

tb = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.2))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "TATA AIA LIFE INSURANCE • THANE / MUMBAI"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_CYAN
p.space_after = Pt(10)

p2 = tf.add_paragraph()
p2.text = "6-Month Strategic Roadmap: AVP - Cloud AI Ops"
p2.font.size = Pt(34)
p2.font.bold = True
p2.font.color.rgb = WHITE
p2.space_after = Pt(14)

p3 = tf.add_paragraph()
p3.text = "Comprehensive Preparation Strategy Aligning Azure Delta Lake, Big Data CDC, Cognitive AI Ops, Enterprise CI/CD & Life Insurance Architecture Governance"
p3.font.size = Pt(15)
p3.font.color.rgb = RGBColor(200, 215, 235)

tb_foot = slide1.shapes.add_textbox(Inches(1.2), Inches(6.2), Inches(11.0), Inches(0.8))
tf_f = tb_foot.text_frame
p_f = tf_f.paragraphs[0]
p_f.text = "Core Pillars: Azure Cloud Architecture | AI Ops | Delta Lake & Big Data | DevOps CI/CD | Risk & IRDAI Governance"
p_f.font.size = Pt(12)
p_f.font.color.rgb = GOLD

# SLIDE 2: Role & JD Mapping
create_card_slide(
    prs,
    "Executive Role Mapping & High-Level JD Mandate",
    {
        "title": "📋 Key Leadership Responsibilities",
        "items": [
            "Align business solution architecture with Alliance standards & customer requirements.",
            "Architect reusable Service Catalog components to accelerate enterprise delivery.",
            "Drive CI/CD DevOps toolchain implementation across hybrid cloud environments.",
            "Conduct capacity planning, technical cost/performance trade-off evaluations.",
            "Act as technical liaison between Business Units & IT / Data Center delivery teams.",
            "Manage enterprise IT architecture risks, IRDAI compliance & security controls."
        ]
    },
    {
        "title": "⚙️ Mandatory Technical Expertise",
        "items": [
            "7+ Years Azure IaaS, Monitoring, Cloud Backup & Site Recovery Services.",
            "Must-Have Core Skill: Azure Delta Lake, Big Data Architecture & CDC (Change Data Capture).",
            "Cognitive AI Services, Azure OpenAI, MLOps & Intelligent Telemetry Automation.",
            "Azure Stack Hub / HCI Hybrid Cloud Infrastructure & Hyper-V Virtual Networking.",
            "Azure Sentinel, Azure AD, MECM (SCCM) & Update Management.",
            "Mentoring Java / J2EE / SOA application teams to Modern Cloud Native Microservices."
        ]
    }
)

# SLIDE 3: Month 1
create_card_slide(
    prs,
    "Month 1: Azure Enterprise Foundation & Hybrid Cloud",
    {
        "title": "🎯 Learning Objectives & Key Topics",
        "items": [
            "Master Azure Core IaaS/PaaS architecture & Enterprise Landing Zones.",
            "Deep dive into Azure Stack Hub & Azure Stack HCI hybrid deployment patterns.",
            "Understand Hyper-V virtual networking, Azure AD tenant structures & RBAC.",
            "Explore Azure Backup & Azure Site Recovery (ASR) disaster recovery topologies.",
            "Analyze Tata AIA Life Insurance core domain: Policy Admin, Claims & Tele-underwriting."
        ]
    },
    {
        "title": "🚀 Action Plan & Hands-on Deliverables",
        "items": [
            "Build a hybrid connectivity POC connecting on-prem Hyper-V to Azure Virtual Network.",
            "Configure Azure Site Recovery (ASR) vault with automated failover testing script.",
            "Draft Enterprise Azure Landing Zone architecture diagram adhering to security standards.",
            "Implement Azure Update Management & MECM hybrid patch management policy.",
            "Milestone: Complete Azure Solutions Architect (AZ-305) hybrid refresher."
        ]
    }
)

# SLIDE 4: Month 2
create_card_slide(
    prs,
    "Month 2: Azure Delta Lake, Big Data & CDC Mastery",
    {
        "title": "🎯 Learning Objectives & Key Topics",
        "items": [
            "Master Medallion Architecture in Azure Delta Lake (Bronze -> Silver -> Gold).",
            "Implement Change Data Capture (CDC) using Qlik / Debezium / Azure Data Factory.",
            "Azure Databricks & PySpark optimization for real-time telemetry & policy logs.",
            "Azure SQL Database, Synapse Analytics & ADLS Gen2 storage lifecycle rules.",
            "Big Data capacity management & cost optimization techniques for insurance telemetry."
        ]
    },
    {
        "title": "🚀 Action Plan & Hands-on Deliverables",
        "items": [
            "Construct a end-to-end CDC pipeline capturing SQL Server transactions into Delta Lake.",
            "Optimize PySpark/Delta lake query performance with Z-Ordering & Liquid Clustering.",
            "Build real-time streaming ingest pipeline for insurance claims & tele-underwriting logs.",
            "Design cost tiering model for ADLS Gen2 archiving (Hot -> Cool -> Cold -> Archive).",
            "Milestone: Azure Data Engineer (DP-203) & Delta Lake Architecture certification."
        ]
    }
)

# SLIDE 5: Month 3
create_card_slide(
    prs,
    "Month 3: Cloud AI Ops, Cognitive Services & Sentinel",
    {
        "title": "🎯 Learning Objectives & Key Topics",
        "items": [
            "Azure Cognitive Services (Document Intelligence, Vision, Language, AI Search).",
            "Azure OpenAI & LLM integration for automated claim document ingestion.",
            "AI Ops: Automated anomaly detection in application telemetry via Azure Monitor & Log Analytics.",
            "Azure Sentinel SIEM integration for security event monitoring & playbook automation.",
            "MLOps pipelines: Model registration, deployment & drift monitoring in Azure Machine Learning."
        ]
    },
    {
        "title": "🚀 Action Plan & Hands-on Deliverables",
        "items": [
            "Build AI Ops pipeline auto-detecting latency anomalies & triggering self-healing webhooks.",
            "Develop POC for automated OCR & claim form classification using Azure Document Intelligence.",
            "Deploy Azure Sentinel incident response playbook using Azure Logic Apps.",
            "Establish automated capacity prediction model for cloud server scaling.",
            "Milestone: Complete Azure AI Engineer (AI-102) & AI Ops Playbook blueprint."
        ]
    }
)

# SLIDE 6: Month 4
create_card_slide(
    prs,
    "Month 4: Enterprise DevOps, CI/CD & Service Catalog",
    {
        "title": "🎯 Learning Objectives & Key Topics",
        "items": [
            "Azure DevOps Pipelines (YAML), GitHub Actions & Infrastructure as Code (Terraform/Bicep).",
            "Reusable Service Catalog development: Modular Bicep/Terraform templates for teams.",
            "Integration Testing & Automated Release Quality Gates in CI/CD pipelines.",
            "Configuration Management DB (CMDB) & Azure Resource Graph tracking.",
            "Capacity Planning & Chaos Engineering resilience testing."
        ]
    },
    {
        "title": "🚀 Action Plan & Hands-on Deliverables",
        "items": [
            "Build an Enterprise Service Catalog repository containing approved IaC modules.",
            "Create end-to-end multi-stage YAML pipeline with automated security scan (Checkov/SonarQube).",
            "Implement automated integration testing framework for cloud service deployments.",
            "Establish automated capacity alerts & auto-scaling policies for peak insurance periods.",
            "Milestone: Complete DevOps Engineer Expert (AZ-400) & Service Catalog launch."
        ]
    }
)

# SLIDE 7: Month 5
create_card_slide(
    prs,
    "Month 5: Insurance Architecture, Security & IRDAI",
    {
        "title": "🎯 Learning Objectives & Key Topics",
        "items": [
            "IRDAI Cloud Security Guidelines & Data Residency requirements for Life Insurance.",
            "Enterprise Security Architecture: Azure Sentinel, Key Vault, Private Endpoints & WAF.",
            "Risk mitigation frameworks for non-stated functional requirements (availability, latency).",
            "Trade-off analysis methodology: Cost vs Scalability vs Security in cloud design.",
            "Coordination workflows between IT, Data Center (DC), Cloud & Business Verticals."
        ]
    },
    {
        "title": "🚀 Action Plan & Hands-on Deliverables",
        "items": [
            "Draft IRDAI Cloud Compliance Audit Checklist covering data encryption & sovereignty.",
            "Author a formal Solution Architecture Document (SAD) including trade-off matrices.",
            "Implement Zero-Trust network topology using Azure Private Link & Firewall Premium.",
            "Conduct disaster recovery & business continuity drill simulation (RTO < 15 mins, RPO < 5 mins).",
            "Milestone: Complete Cybersecurity Architect (SC-100) & Risk Governance Framework."
        ]
    }
)

# SLIDE 8: Month 6
create_card_slide(
    prs,
    "Month 6: Executive Leadership & Team Mentorship",
    {
        "title": "🎯 Learning Objectives & Key Topics",
        "items": [
            "Mentoring legacy Java/J2EE/SOA teams toward Cloud-Native Spring Boot / Microservices.",
            "Architecture Review Board (ARB) presentation & stakeholder management skills.",
            "Sequencing project deliverables to prevent rework & fit building blocks together.",
            "Vendor & Alliance relationship management, effort estimation & team capacity planning.",
            "30-60-90 Day Execution Vision on taking charge at Tata AIA Thane."
        ]
    },
    {
        "title": "🚀 Action Plan & Hands-on Deliverables",
        "items": [
            "Develop 'Java/SOA to Azure Microservices' Migration Training Curriculum for developers.",
            "Establish Architecture Review Board (ARB) submission template & verification checklist.",
            "Create effort estimation model (T-Shirt sizing & Story Point calibration) for cloud projects.",
            "Conduct mock executive presentation for cloud transformation strategy.",
            "Milestone: Complete Leadership Readiness & Mock Executive ARB Defense."
        ]
    }
)

# SLIDE 9: Technical Mastery Checklist
slide_matrix = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide_matrix, CREAM)
add_header(slide_matrix, "6-Month Technical & Competency Mastery Checklist")

rows, cols = 6, 4
left, top, width, height = Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.4)
table_shape = slide_matrix.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table

table.columns[0].width = Inches(2.2)
table.columns[1].width = Inches(3.4)
table.columns[2].width = Inches(4.3)
table.columns[3].width = Inches(1.8)

headers = ["Domain Area", "Required Skill in JD", "Target Mastery Output", "Status Target"]
for i, h in enumerate(headers):
    cell = table.cell(0, i)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = DARK_NAVY
    p = cell.text_frame.paragraphs[0]
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.font.size = Pt(12)

matrix_data = [
    ("Delta Lake & CDC", "Azure Delta Lake, Big Data, CDC, Synapse", "Real-time streaming CDC pipeline to Medallion Delta Lake", "Month 2 Complete"),
    ("Cloud AI Ops", "Cognitive Services, AI/ML, AI Ops monitoring", "AI Ops auto-healing anomaly detection & Document AI", "Month 3 Complete"),
    ("DevOps & Catalog", "Azure DevOps, IaC, Reusable Service Catalog", "Enterprise Bicep Service Catalog & CI/CD pipeline", "Month 4 Complete"),
    ("Security & IRDAI", "Azure Sentinel, Hybrid AD, IRDAI Compliance", "IRDAI Compliant Zero-Trust Cloud Architecture SAD", "Month 5 Complete"),
    ("Mentorship & ARB", "Java/SOA Mentoring, ARB, Effort Estimation", "Modernization curriculum & Effort Estimation Tool", "Month 6 Complete")
]

for row_idx, data in enumerate(matrix_data, start=1):
    for col_idx, text in enumerate(data):
        cell = table.cell(row_idx, col_idx)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE if row_idx % 2 == 1 else RGBColor(240, 243, 246)
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        if col_idx == 3:
            p.font.bold = True
            p.font.color.rgb = GREEN

# SLIDE 10: 30-60-90 Day Plan on Joining
create_card_slide(
    prs,
    "30-60-90 Day Execution Vision Upon Joining Tata AIA",
    {
        "title": "Phase 1: Days 1 to 30 (Assessment & Audit)",
        "items": [
            "Audit existing Azure infrastructure, Delta Lake CDC pipelines & Azure Sentinel setup.",
            "Meet key Business & DC/IT stakeholders to identify architectural pain points.",
            "Review existing Java/SOA applications and cloud migration backlog.",
            "Evaluate current DevOps pipeline security and Service Catalog usage.",
            "Deliverable: Baseline Cloud & AI Ops Maturity Assessment Report."
        ]
    },
    {
        "title": "Phase 2 & 3: Days 31 to 90 (Optimization & Scale)",
        "items": [
            "Days 31-60: Standardize Bicep/Terraform Service Catalog and launch AI Ops pilot.",
            "Days 31-60: Implement automated CDC pipeline optimizations for Delta Lake data ingest.",
            "Days 61-90: Establish Architecture Review Board (ARB) governance & IRDAI compliance gate.",
            "Days 61-90: Roll out developer mentorship program for Java to Cloud Native microservices.",
            "Deliverable: Fully Operational Cloud AI Ops Framework at Tata AIA."
        ]
    }
)

# SLIDE 11: Conclusion
slide_end = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide_end, DARK_NAVY)

tb_end = slide_end.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.5))
tf_e = tb_end.text_frame
tf_e.word_wrap = True

p = tf_e.paragraphs[0]
p.text = "TATA AIA LIFE INSURANCE • AVP CLOUD AI OPS"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_CYAN
p.space_after = Pt(10)

p2 = tf_e.add_paragraph()
p2.text = "Roadmap Executed for Technical & Executive Success"
p2.font.size = Pt(32)
p2.font.bold = True
p2.font.color.rgb = WHITE
p2.space_after = Pt(16)

p3 = tf_e.add_paragraph()
p3.text = "By following this structured 6-month plan, you will achieve complete mastery over Azure Delta Lake CDC, Cognitive AI Ops, Enterprise CI/CD, and Life Insurance Architecture Governance — positioning you as the ideal AVP leader for Tata AIA Thane."
p3.font.size = Pt(15)
p3.font.color.rgb = RGBColor(200, 215, 235)

output_path = "/Users/neeraj.jha/forbes-fab-luxe/Tata_AIA_AVP_Cloud_AIOps_6Month_Plan.pptx"
prs.save(output_path)
print(f"Complete PPTX saved successfully to {output_path}")
