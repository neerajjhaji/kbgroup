import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Tata AIA Brand Palette
DARK_NAVY = RGBColor(11, 25, 44)       # #0B192C Tata AIA Navy
TATA_BLUE = RGBColor(0, 102, 204)      # #0066CC Tata Blue
ACCENT_CYAN = RGBColor(0, 180, 216)    # #00B4D8 AI/Cloud Cyan
CREAM = RGBColor(248, 249, 250)        # #F8F9FA Off-White
CARD_BG = RGBColor(255, 255, 255)      # White Card
TEXT_DARK = RGBColor(33, 37, 41)
TEXT_MUTED = RGBColor(108, 117, 125)
WHITE = RGBColor(255, 255, 255)
GOLD = RGBColor(212, 175, 55)
GREEN = RGBColor(40, 167, 69)

def set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def add_header(slide, title_text, category_text="TATA AIA LIFE INSURANCE • AVP CLOUD AI OPS PROBATION CONFIRMATION PLAN"):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = TATA_BLUE
    p0.space_after = Pt(2)

    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = DARK_NAVY

def create_card_slide(prs, title, card1_data, card2_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, CREAM)
    add_header(slide, title)

    top_pos = Inches(1.4)
    card_width = Inches(5.6)
    card_height = Inches(5.4)

    # Left Card
    c1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, card_width, card_height)
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = TATA_BLUE
    c1.line.width = Pt(1.5)

    tf1 = c1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = card1_data['title']
    p1.font.bold = True
    p1.font.size = Pt(15)
    p1.font.color.rgb = TATA_BLUE
    p1.space_after = Pt(10)

    for item in card1_data['items']:
        p = tf1.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # Right Card
    c2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), top_pos, card_width, card_height)
    c2.fill.solid()
    c2.fill.fore_color.rgb = DARK_NAVY
    c2.line.color.rgb = ACCENT_CYAN
    c2.line.width = Pt(1.5)

    tf2 = c2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = card2_data['title']
    p2.font.bold = True
    p2.font.size = Pt(15)
    p2.font.color.rgb = ACCENT_CYAN
    p2.space_after = Pt(10)

    for item in card2_data['items']:
        p = tf2.add_paragraph()
        p.text = f"➔ {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = WHITE
        p.space_after = Pt(6)

# SLIDE 1: Title Slide
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide1, DARK_NAVY)

bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.08), Inches(4.0))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT_CYAN
bar.line.fill.background()

tb = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.2))
tf = tb.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "TATA AIA LIFE INSURANCE • THANE / MUMBAI"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_CYAN
p.space_after = Pt(10)

p2 = tf.add_paragraph()
p2.text = "30-60-90 Day Probation & Confirmation Plan"
p2.font.size = Pt(34)
p2.font.bold = True
p2.font.color.rgb = WHITE
p2.space_after = Pt(14)

p3 = tf.add_paragraph()
p3.text = "Strategic Roadmap to Demonstrate Leadership, Technical Depth, AI Ops Governance, and Secure Executive Confirmation as AVP - Cloud AI Ops"
p3.font.size = Pt(15)
p3.font.color.rgb = RGBColor(200, 215, 235)

tb_foot = slide1.shapes.add_textbox(Inches(1.2), Inches(6.2), Inches(11.0), Inches(0.8))
tf_f = tb_foot.text_frame
p_f = tf_f.paragraphs[0]
p_f.text = "Focus: Audit & Quick Wins (Days 1-30) | Standardize & Pilot (Days 31-60) | Scale & Confirm (Days 61-90)"
p_f.font.size = Pt(12)
p_f.font.color.rgb = GOLD

# SLIDE 2: Executive Summary & Objective
create_card_slide(
    prs,
    "Executive Overview & Probation Success Objectives",
    {
        "title": "🎯 Strategic Probation Mandate",
        "items": [
            "Seamlessly take ownership of Tata AIA Cloud & AI Ops estate within 90 days.",
            "Establish immediate credibility with Business Verticals, DC/IT Teams & Alliance Partners.",
            "Eliminate cloud architecture bottlenecks in Delta Lake, CDC, and DevOps pipelines.",
            "Enforce IRDAI insurance compliance, zero-trust security & risk management.",
            "Demonstrate AVP-level leadership through developer mentorship & ARB governance."
        ]
    },
    {
        "title": "🏆 Key Confirmation Milestones",
        "items": [
            "Day 30: State-of-Cloud Audit Report & 2 High-Impact Quick Wins delivered.",
            "Day 60: Enterprise Reusable Service Catalog & AI Ops Anomaly Pilot live.",
            "Day 90: IRDAI Zero-Trust Compliance Blueprint & ARB Governance framework enforced.",
            "Day 90: Java/SOA to Cloud-Native Mentorship curriculum fully deployed.",
            "Result: 100% Executive Endorsement for Formal Role Confirmation."
        ]
    }
)

# SLIDE 3: Days 1-30 (Phase 1)
create_card_slide(
    prs,
    "Days 1 – 30 | Phase 1: Audit, Stakeholder Alignment & Quick Wins",
    {
        "title": "🔍 Discovery & Infrastructure Audit",
        "items": [
            "Audit Azure footprint: Landing zones, Virtual Networks, Azure Stack HCI & Hyper-V.",
            "Inspect Delta Lake CDC pipelines, ADLS storage tiers, and PySpark query latency.",
            "Review Azure Sentinel SIEM, Log Analytics workspaces, and Azure AD RBAC policies.",
            "Assess current Azure DevOps pipelines, Bicep/Terraform scripts, and CMDB assets.",
            "Map core life insurance workflows: Policy Admin, Tele-underwriting, and Claims."
        ]
    },
    {
        "title": "⚡ Stakeholder Engagement & Quick Wins",
        "items": [
            "Conduct 1-on-1 interviews with Head of DC, EA, InfoSec, and Business Unit leads.",
            "Quick Win 1: Optimize a slow CDC data ingestion pipeline for immediate speedup.",
            "Quick Win 2: Clean up orphaned Azure resources for instant monthly cost savings.",
            "Establish weekly cadence between Cloud Architecture and Data Center teams.",
            "Deliverable: Day 30 Comprehensive Cloud Maturity Assessment Report."
        ]
    }
)

# SLIDE 4: Days 31-60 (Phase 2)
create_card_slide(
    prs,
    "Days 31 – 60 | Phase 2: Standardize, Pilot & AI Ops Automation",
    {
        "title": "⚙️ Service Catalog & AI Ops Pilot",
        "items": [
            "Publish Enterprise Reusable Service Catalog with approved Bicep/Terraform modules.",
            "Deploy AI Ops Anomaly Detection pilot in Azure Monitor to auto-detect latency spikes.",
            "Optimize Delta Lake Medallion Architecture (Bronze -> Silver -> Gold) with Liquid Clustering.",
            "Implement automated quality gates in Azure DevOps YAML pipelines (Checkov/SonarQube).",
            "Establish capacity management baseline for peak policy renewal traffic."
        ]
    },
    {
        "title": "🎓 Mentorship & Governance Framework",
        "items": [
            "Launch 'Java/SOA to Cloud Native' mentorship cohort for application development teams.",
            "Draft formal IRDAI Cloud Security & Data Sovereignty Audit Checklist.",
            "Implement Zero-Trust network rules using Azure Private Link & Key Vault.",
            "Establish T-Shirt effort estimation model for new business solution requests.",
            "Deliverable: Day 60 Operational Service Catalog & AI Ops Pilot Launch."
        ]
    }
)

# SLIDE 5: Days 61-90 (Phase 3)
create_card_slide(
    prs,
    "Days 61 – 90 | Phase 3: Scale, Govern & Confirmation Defense",
    {
        "title": "🚀 Enterprise Scale & DR Readiness",
        "items": [
            "Scale AI Ops auto-healing webhooks across production insurance workloads.",
            "Enforce Architecture Review Board (ARB) design verification gates for all releases.",
            "Conduct full Disaster Recovery (DR) & Business Continuity drill (RTO < 15m, RPO < 5m).",
            "Finalize capacity planning and auto-scaling rules for business growth.",
            "Verify all IT design and cloud architecture risks are logged and mitigated."
        ]
    },
    {
        "title": "🏆 Confirmation Presentation & Sign-off",
        "items": [
            "Compile Executive Probation Achievement Report demonstrating tangible ROI & speed.",
            "Highlight developer capability uplift from Java/SOA to Cloud Microservices.",
            "Present formal 90-Day Transformation Review to CIO, CTO & IT Leadership.",
            "Obtain unanimous Architecture Review Board & HR confirmation endorsement.",
            "Deliverable: Formal Confirmation as AVP - Cloud AI Ops, Tata AIA."
        ]
    }
)

# SLIDE 6: Skill & Impact Matrix
slide_matrix = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide_matrix, CREAM)
add_header(slide_matrix, "Probation Confirmation Value Matrix & KPI Indicators")

rows, cols = 6, 4
left, top, width, height = Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.4)
table_shape = slide_matrix.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table

table.columns[0].width = Inches(2.2)
table.columns[1].width = Inches(3.4)
table.columns[2].width = Inches(4.3)
table.columns[3].width = Inches(1.8)

headers = ["Probation Phase", "Strategic Focus Area", "Measurable Key Result (KPI)", "Confirmation Impact"]
for i, h in enumerate(headers):
    cell = table.cell(0, i)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = DARK_NAVY
    p = cell.text_frame.paragraphs[0]
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.font.size = Pt(12)

matrix_data = [
    ("Days 1 - 30", "Audit & Quick Wins", "2 Quick wins delivered; 100% audit of Delta Lake & CDC pipelines", "Immediate Credibility"),
    ("Days 31 - 60", "Service Catalog & AI Ops", "10+ Reusable IaC modules; AI Ops anomaly auto-alerts live", "Architectural Speed"),
    ("Days 31 - 60", "Team Mentorship", "Java/SOA developers trained on Cloud Native Bicep & Microservices", "Team Capability Uplift"),
    ("Days 61 - 90", "IRDAI Security & ARB", "Zero-Trust network, DR drill passed (RTO < 15m), ARB process live", "Enterprise Governance"),
    ("Day 90", "Executive Presentation", "100% Endorsement from CIO / IT Leadership for formal confirmation", "CONFIRMED AVP")
]

for row_idx, data in enumerate(matrix_data, start=1):
    for col_idx, text in enumerate(data):
        cell = table.cell(row_idx, col_idx)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE if row_idx % 2 == 1 else RGBColor(240, 243, 246)
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        if col_idx == 3:
            p.font.bold = True
            p.font.color.rgb = GREEN

# SLIDE 7: Conclusion
slide_end = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide_end, DARK_NAVY)

tb_end = slide_end.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.5))
tf_e = tb_end.text_frame
tf_e.word_wrap = True

p = tf_e.paragraphs[0]
p.text = "TATA AIA LIFE INSURANCE • PROBATION CONFIRMATION ROADMAP"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_CYAN
p.space_after = Pt(10)

p2 = tf_e.add_paragraph()
p2.text = "Guaranteed Execution for Early Confirmation"
p2.font.size = Pt(32)
p2.font.bold = True
p2.font.color.rgb = WHITE
p2.space_after = Pt(16)

p3 = tf_e.add_paragraph()
p3.text = "By executing this structured 30-60-90 Day Plan, you will demonstrate immediate value, solve critical cloud/data bottlenecks, mentor engineering teams, and deliver IRDAI-compliant AI Ops architecture — ensuring effortless confirmation as AVP - Cloud AI Ops at Tata AIA."
p3.font.size = Pt(15)
p3.font.color.rgb = RGBColor(200, 215, 235)

output_path = "/Users/neeraj.jha/forbes-fab-luxe/Tata_AIA_AVP_30_60_90_Day_Probation_Plan.pptx"
prs.save(output_path)
print(f"Probation PPTX saved successfully to {output_path}")
